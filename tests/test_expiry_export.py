import io
import zipfile
from datetime import date, timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.utils import timezone
from openpyxl import load_workbook

from apps.catalog.models import Platform, Skill
from apps.community.models import ModerationLog
from apps.credentials import expiry, export
from apps.credentials import services as creds
from apps.gamification import game
from apps.notifications.models import Notification

pytestmark = pytest.mark.django_db
PDF = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n%%EOF\n" + b"x" * 200
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


@pytest.fixture(autouse=True)
def seeded(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"
    call_command("seed_catalog", verbosity=0)


def make(owner, lead=None, *, verify=True, data=PDF, name="c.pdf", **extra):
    fields = {"title": "Curso", "kind": "completion", "visibility": "team", **extra}
    cred = creds.create_credential(owner, fields, SimpleUploadedFile(name, data))
    if verify:
        creds.verify(cred, lead)
    return cred


def today():
    return timezone.localdate()


def notices(person, kind):
    return Notification.objects.filter(recipient=person, kind=kind)


# --- vencimientos ----------------------------------------------------------------------------------------------------


def test_expiring_and_expired_sets(member, lead):
    soon = make(member, lead, title="Pronto", expires_on=today() + timedelta(days=10))
    far = make(member, lead, title="Lejos", expires_on=today() + timedelta(days=500))
    gone = make(member, lead, title="Vencida", expires_on=today() - timedelta(days=1))
    pending = make(
        member, lead, verify=False, title="Pendiente", expires_on=today() + timedelta(days=3)
    )
    assert list(expiry.expiring()) == [soon]
    assert list(expiry.expired()) == [gone]
    assert far and pending
    edge = make(member, lead, title="Hoy", expires_on=today())
    assert edge in expiry.expiring() and edge not in expiry.expired()


def test_notify_all_tells_owner_and_leads_once(member, lead):
    make(member, lead, title="Pronto", expires_on=today() + timedelta(days=10))
    make(member, lead, title="Vencida", expires_on=today() - timedelta(days=2))
    assert expiry.notify_all()[:2] == (1, 1)
    assert notices(member, "credential_expiring").count() == 1
    assert notices(member, "credential_expired").count() == 1
    assert (
        notices(lead, "credential_expiring").count() == 1
        and notices(lead, "credential_expired").count() == 1
    )
    expiry.notify_all()
    assert (
        Notification.objects.filter(kind__in=["credential_expiring", "credential_expired"]).count()
        == 4
    )


def test_command_runs(member, lead, capsys):
    make(member, lead, expires_on=today() + timedelta(days=1))
    call_command("check_expirations")
    assert "Por vencer: 1" in capsys.readouterr().out


def test_wings_turn_off_when_the_license_expires_but_xp_is_kept(member, lead):
    platform = Platform.objects.get(slug="dgac-chile")
    cred = make(
        member, lead, kind="license", platform=platform, expires_on=today() + timedelta(days=5)
    )
    assert "alas-dgac" in game.unlocked_badges(member)
    xp = game.total_xp(member)
    cred.expires_on = today() - timedelta(days=1)
    cred.save(update_fields=["expires_on"])
    assert expiry.notify_all()[2] == 1
    assert "alas-dgac" not in game.unlocked_badges(member)
    assert game.total_xp(member) == xp


def test_expirations_page_scopes(client_for, member, lead, admin_person):
    make(member, lead, title="Curso de Ana", expires_on=today() + timedelta(days=5))
    make(lead, admin_person, title="Curso del lead", expires_on=today() + timedelta(days=6))

    def main(client, url):
        return client.get(url).content.decode().split("<main")[1]

    mine = main(client_for(member.login), "/certificados/vencimientos/")
    assert "Curso de Ana" in mine and "Curso del lead" not in mine
    # un miembro no amplía su alcance con ?equipo=1
    assert "Curso del lead" not in main(
        client_for(member.login), "/certificados/vencimientos/?equipo=1"
    )
    lead_own = main(client_for(lead.login), "/certificados/vencimientos/")
    assert "Curso del lead" in lead_own and "Curso de Ana" not in lead_own
    team = main(client_for(lead.login), "/certificados/vencimientos/?equipo=1")
    assert "Curso de Ana" in team and "Curso del lead" in team


def test_home_counter_links_to_expirations(client_for, member, lead):
    make(member, lead, expires_on=today() + timedelta(days=5))
    assert (
        'href="/certificados/vencimientos/"' in client_for(member.login).get("/").content.decode()
    )


# --- exportación -------------------------------------------------------------------------------------------------------


def open_zip(data):
    zf = zipfile.ZipFile(io.BytesIO(data))
    return zf, load_workbook(io.BytesIO(zf.read("evidencia.xlsx")))


def test_only_leads_export(client_for, member, lead):
    make(member, lead)
    assert client_for(member.login).get("/certificados/exportar/").status_code == 403
    assert client_for(member.login).post("/certificados/exportar/", {}).status_code == 403
    with pytest.raises(PermissionError):
        export.build_zip(member, export.select())
    assert client_for(lead.login).get("/certificados/exportar/").status_code == 200


def test_zip_matches_the_selection_and_the_xlsx_is_typed(member, lead, make_person):
    other = make_person("otra@lev.cl")
    make(
        member,
        lead,
        title="Revit básico",
        issuer="Autodesk",
        credential_id="ID-1",
        issued_on=date(2026, 3, 4),
        expires_on=date(2028, 3, 4),
        verify_url="https://v.example/1",
    )
    make(other, lead, title="Mapa", data=PNG, name="m.png")
    make(member, lead, verify=False, title="Pendiente")
    data, count = export.build_zip(lead, export.select())
    assert count == 2
    zf, wb = open_zip(data)
    names = set(zf.namelist())
    assert len(names) == 3 and "evidencia.xlsx" in names
    assert sorted(n.rsplit(".", 1)[1] for n in names if n != "evidencia.xlsx") == ["pdf", "png"]
    ws = wb["Credenciales"]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["Persona", "Credencial"] and "SHA-256" in header
    rows = {r[1].value: r for r in ws.iter_rows(min_row=2)}
    row = rows["Revit básico"]
    assert row[5].value.date() == date(2026, 3, 4) and row[6].value.date() == date(2028, 3, 4)
    assert row[5].number_format == "yyyy-mm-dd"
    assert row[4].value == "ID-1" and row[8].value == "Verificada"
    stored = row[12].value
    assert stored in names and zf.read(stored).startswith(b"%PDF")
    assert rows["Mapa"][5].value is None


def test_filters(member, lead, make_person):
    other = make_person("otra@lev.cl")
    platform = Platform.objects.get(slug="autodesk-learning")
    skill = Skill.objects.get(slug="lidar")
    a = make(member, lead, title="A", platform=platform, kind="certification")
    a.skills.add(skill)
    b = make(other, lead, title="B")
    c = make(member, lead, verify=False, title="C")
    titles = lambda **kw: {x.title for x in export.select(**kw)}  # noqa: E731
    assert titles() == {"A", "B"}
    assert titles(only_verified=False) == {"A", "B", "C"}
    assert titles(people=[other.pk]) == {"B"}
    assert titles(platforms=["autodesk-learning"]) == {"A"}
    assert titles(kinds=["certification"]) == {"A"}
    assert titles(skills=["lidar"]) == {"A"}
    assert titles(kinds=["nope"]) == set()  # un tipo desconocido no amplía la selección
    assert b and c


def test_formula_injection_is_neutralised(member, lead):
    make(member, lead, title='=HYPERLINK("http://evil","x")', issuer="+cmd|' /C calc'!A0")
    _, wb = open_zip(export.build_zip(lead, export.select())[0])
    row = next(wb["Credenciales"].iter_rows(min_row=2))
    assert row[1].value.startswith("'=") and row[3].value.startswith("'+")
    assert row[1].data_type == "s"


def test_zip_names_are_generated_and_unique(member, lead):
    make(member, lead, title="../../etc/passwd")
    make(member, lead, title="../../etc/passwd")
    zf, _ = open_zip(export.build_zip(lead, export.select())[0])
    names = [n for n in zf.namelist() if n != "evidencia.xlsx"]
    assert len(set(names)) == 2 and all(".." not in n and not n.startswith("/") for n in names)


def test_missing_file_is_reported_not_fatal(member, lead):
    import os

    cred = make(member, lead, title="Sin archivo")
    os.remove(cred.file.path)
    zf, wb = open_zip(export.build_zip(lead, export.select())[0])
    assert next(wb["Credenciales"].iter_rows(min_row=2))[12].value == "(archivo no disponible)"
    assert zf.namelist() == ["evidencia.xlsx"]


def test_empty_and_oversized_selections(member, lead, monkeypatch):
    with pytest.raises(ValueError, match="No hay"):
        export.build_zip(lead, export.select())
    make(member, lead)
    monkeypatch.setattr(export, "MAX_CREDENTIALS", 0)
    with pytest.raises(ValueError, match="supera"):
        export.build_zip(lead, export.select())
    monkeypatch.setattr(export, "MAX_CREDENTIALS", 500)
    monkeypatch.setattr(export, "MAX_BYTES", 10)
    with pytest.raises(ValueError, match="tamaño"):
        export.build_zip(lead, export.select())


def test_http_download_and_audit_log(client_for, member, lead):
    make(member, lead, title="Exportable")
    client = client_for(lead.login)
    preview = client.get("/certificados/exportar/").content.decode()
    assert "Exportable" in preview and "Descargar ZIP" in preview
    response = client.post("/certificados/exportar/", {"estado": "verified"})
    assert response.status_code == 200 and response["Content-Type"] == "application/zip"
    assert (
        "attachment" in response["Content-Disposition"] and response["Cache-Control"] == "no-store"
    )
    zf, _ = open_zip(response.content)
    assert len(zf.namelist()) == 2
    assert ModerationLog.objects.filter(
        action="export_credentials", actor=lead, summary="1 credenciales"
    ).exists()


def test_http_filtered_download(client_for, member, lead, make_person):
    other = make_person("otra@lev.cl")
    make(member, lead, title="De ana")
    make(other, lead, title="De otra")
    response = client_for(lead.login).post(
        "/certificados/exportar/", {"persona": [other.pk], "estado": "verified"}
    )
    _, wb = open_zip(response.content)
    assert [r[1].value for r in wb["Credenciales"].iter_rows(min_row=2)] == ["De otra"]


def test_empty_http_selection_shows_a_message_and_logs_nothing(client_for, lead):
    response = client_for(lead.login).post(
        "/certificados/exportar/", {"estado": "verified"}, follow=True
    )
    assert "No hay credenciales" in response.content.decode()
    assert not ModerationLog.objects.filter(action="export_credentials").exists()


def test_export_links_only_for_leads(client_for, member, lead):
    assert (
        "Exportar evidencia" not in client_for(member.login).get("/certificados/").content.decode()
    )
    assert "Exportar evidencia" in client_for(lead.login).get("/certificados/").content.decode()
